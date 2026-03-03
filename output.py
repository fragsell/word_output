import zipfile
import os
import json
import re
from docx import Document


def parse_docx_by_section_titles(file_path):
    """Парсит документ по названиям разделов"""
    doc = Document(file_path)

    # Собираем все параграфы
    all_paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            all_paragraphs.append(text)

    # Заголовок документа
    title = all_paragraphs[0] if all_paragraphs else ""

    # Список названий разделов в правильном порядке
    section_titles = [
        "Область применения",
        "Соответствие стандартам и спецификации",
        "Основные параметры",
        "Материалы исполнения",
        "Рабочие параметры",
        "Конструктивные особенности",
        "Маркировка",
        "Контроль качества",
        "Упаковка и хранение",
        "Гарантия",
        "Сведения о поставке"
    ]

    # Создаем структуру для хранения данных
    sections = {}
    current_section = None
    section_content = []

    # Проходим по всем параграфам, начиная с "ТЕХНИЧЕСКИЙ ПАСПОРТ"
    start_index = 0
    for i, p in enumerate(all_paragraphs):
        if "ТЕХНИЧЕСКИЙ ПАСПОРТ" in p:
            start_index = i + 1
            break

    # Собираем данные по разделам
    for i in range(start_index, len(all_paragraphs)):
        p = all_paragraphs[i]

        # Проверяем, является ли это названием раздела
        is_section_title = False
        section_name = None

        for title in section_titles:
            if p == title:
                is_section_title = True
                section_name = title
                break

        if is_section_title:
            # Сохраняем предыдущий раздел
            if current_section is not None and section_content:
                section_text = "\n".join(section_content).strip()
                sections[current_section] = section_text

            # Начинаем новый раздел
            current_section = section_name
            section_content = []
        elif current_section is not None:
            # Добавляем текст к текущему разделу
            section_content.append(p)

    # Сохраняем последний раздел
    if current_section is not None and section_content:
        section_text = "\n".join(section_content).strip()
        sections[current_section] = section_text

    # Извлекаем таблицу
    table_data = extract_table_data(doc)

    return {
        "document_title": all_paragraphs[0] if all_paragraphs else "",
        "sections": sections,
        "table_data": table_data
    }


def extract_table_data(doc):
    """Извлекает данные из таблиц"""
    for table in doc.tables:
        rows = []
        headers = []

        for i, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]

            if i == 0:
                headers = cells
            else:
                row_dict = {}
                for j, cell_text in enumerate(cells):
                    if j < len(headers) and headers[j]:
                        # Очищаем названия заголовков
                        header = headers[j].replace('\n', ' ').strip()
                        row_dict[header] = cell_text

                if row_dict and any(row_dict.values()):
                    rows.append(row_dict)

        if rows and any("DN" in key for key in (headers + list(rows[0].keys()) if rows else [])):
            return rows

    return []


def clean_text(text):
    """Очищает текст"""
    if not text:
        return ""

    # Заменяем переносы строк в середине предложений
    text = re.sub(r'([а-яa-z])\n([а-яa-z])',
                  r'\1 \2', text, flags=re.IGNORECASE)

    # Удаляем лишние пробелы
    text = re.sub(r'\s+', ' ', text)

    return text.strip()


def create_structured_data(data):
    """Создает структурированные данные в нужном формате"""
    
    sections = data.get("sections", {})
    table_data = data.get("table_data", [])
    
    # Формируем описание из трех разделов
    description_parts = []
    
    if "Область применения" in sections:
        description_parts.append(clean_text(sections["Область применения"]))
    
    if "Соответствие стандартам и спецификации" in sections:
        description_parts.append(clean_text(sections["Соответствие стандартам и спецификации"]))
    
    if "Конструктивные особенности" in sections:
        description_parts.append(clean_text(sections["Конструктивные особенности"]))
    
    # Объединяем описание
    description = " ".join(description_parts)
    
    # Рабочие параметры
    working_parameters = clean_text(sections.get("Рабочие параметры", ""))
    
    # Материалы исполнения
    materials = clean_text(sections.get("Материалы исполнения", ""))
    
    # Создаем итоговую структуру
    structured_data = {
        "описание": description,
        "рабочие_параметры": working_parameters,
        "материалы_исполнения": materials,
        "таблица": table_data
    }
    
    return structured_data


def save_to_json(data, output_file):
    """Сохраняет данные в JSON"""
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=4)


def main():
    input_file = "Гайка шлицевая молочная, резьбовая, DIN AISI 316.docx"
    output_file = "гайка_структурированные_данные.json"

    print("=" * 80)
    print("ИЗВЛЕЧЕНИЕ ДАННЫХ ИЗ ДОКУМЕНТА")
    print("=" * 80)

    try:
        # Парсим документ
        raw_data = parse_docx_by_section_titles(input_file)
        
        # Создаем структурированные данные
        structured_data = create_structured_data(raw_data)

        # Выводим результаты
        print(f"\nЗаголовок документа: {raw_data['document_title']}")
        print(f"\nНайдено разделов: {len(raw_data['sections'])}")

        # Список разделов в правильном порядке для вывода
        ordered_sections = [
            "Область применения",
            "Соответствие стандартам и спецификации",
            "Основные параметры",
            "Материалы исполнения",
            "Рабочие параметры",
            "Конструктивные особенности",
            "Маркировка",
            "Контроль качества",
            "Упаковка и хранение",
            "Гарантия",
            "Сведения о поставке"
        ]

        print("\nСОДЕРЖАНИЕ РАЗДЕЛОВ:")
        print("-" * 80)

        for section_title in ordered_sections:
            if section_title in raw_data["sections"]:
                content = raw_data["sections"][section_title]
                
                print(f"\n{section_title}:")
                if content:
                    # Очищаем и форматируем текст
                    clean_content = clean_text(content)
                    char_count = len(clean_content)

                    print(f"  Длина: {char_count} символов")

                    # Показываем превью
                    preview_lines = clean_content.split('. ')
                    if len(preview_lines) > 0:
                        print(f"  Превью: {preview_lines[0][:150]}...")
                else:
                    print("  [Нет содержания]")
            else:
                print(f"\n{section_title}: [Раздел не найден]")

        # Выводим таблицу
        table_data = raw_data.get("table_data", [])
        if table_data:
            print(f"\n" + "=" * 80)
            print(f"ТАБЛИЦА ПАРАМЕТРОВ: {len(table_data)} строк")
            print("-" * 80)

            # Красиво выводим таблицу
            if table_data:
                headers = list(table_data[0].keys())
                print(" | ".join(headers))
                print("-" * 60)

                for row in table_data[:5]:  # Первые 5 строк
                    values = [str(row.get(h, "")) for h in headers]
                    print(" | ".join(values))

                if len(table_data) > 5:
                    print(f"... и еще {len(table_data) - 5} строк")

        # Выводим структурированные данные
        print(f"\n" + "=" * 80)
        print("СТРУКТУРИРОВАННЫЕ ДАННЫЕ:")
        print("-" * 80)
        
        print(f"\nОПИСАНИЕ (из 3 разделов):")
        print(f"  Длина: {len(structured_data['описание'])} символов")
        if structured_data['описание']:
            print(f"  Превью: {structured_data['описание'][:200]}...")
        
        print(f"\nРАБОЧИЕ ПАРАМЕТРЫ:")
        print(f"  Длина: {len(structured_data['рабочие_параметры'])} символов")
        if structured_data['рабочие_параметры']:
            print(f"  Превью: {structured_data['рабочие_параметры'][:200]}...")
        
        print(f"\nМАТЕРИАЛЫ ИСПОЛНЕНИЯ:")
        print(f"  Длина: {len(structured_data['материалы_исполнения'])} символов")
        if structured_data['материалы_исполнения']:
            print(f"  Превью: {structured_data['материалы_исполнения'][:200]}...")
        
        print(f"\nТАБЛИЦА:")
        print(f"  Количество строк: {len(structured_data['таблица'])}")

        # Сохраняем в JSON
        save_to_json(structured_data, output_file)

        print(f"\n" + "=" * 80)
        print(f"✓ Данные сохранены в файл: {output_file}")

        # Показываем пример JSON структуры (без многоточия)
        print(f"\nПРИМЕР JSON СТРУКТУРЫ:")
        print("-" * 80)
        print("""
{
    "описание": "текст из разделов Область применения, Соответствие стандартам и спецификациям, Конструктивные особенности",
    "рабочие_параметры": "текст из раздела Рабочие параметры",
    "материалы_исполнения": "текст из раздела Материалы исполнения",
    "таблица": [
        {
            "DN": "15",
            "Шаг резьбы": "1.5",
            "d, мм": "22.3"
        }
    ]
}
        """.strip())

        # Проверяем, все ли нужные разделы найдены
        required_sections = [
            "Область применения",
            "Соответствие стандартам и спецификации",
            "Конструктивные особенности",
            "Рабочие параметры",
            "Материалы исполнения"
        ]
        
        missing_sections = [
            s for s in required_sections if s not in raw_data["sections"]]
        if missing_sections:
            print(f"\nВНИМАНИЕ: Отсутствуют разделы: {', '.join(missing_sections)}")

    except Exception as e:
        print(f"✗ Ошибка: {str(e)}")
        import traceback
        traceback.print_exc()


if __name__ == "__main__":
    main()